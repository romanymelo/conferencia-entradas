import logging
import shutil
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas

from conferencia.modelo import Categoria, Status
from conferencia.descoberta import Par, identifica, pareia
from conferencia.leitor_dominio import le_dominio
from conferencia.leitor_sefaz import le_sefaz
from main import executa

# Pasta local com os PDFs reais (um par por empresa), nunca versionada (.gitignore).
# Aceita "amostra" ou "amostras": o pareamento e por CNPJ lido de dentro do PDF, entao
# nem o nome da pasta nem o nome dos arquivos importam para a conferencia.
_CANDIDATAS = ("amostra", "amostras")
AMOSTRA = next(
    (Path(__file__).parent / nome for nome in _CANDIDATAS
     if (Path(__file__).parent / nome).is_dir()),
    Path(__file__).parent / _CANDIDATAS[0],
)


# --- PDFs fabricados (dados fictícios) ---

def _pdf(caminho, linhas):
    c = canvas.Canvas(str(caminho), pagesize=landscape(A4))
    y = 550
    for texto in linhas:
        c.drawString(20, y, texto)
        y -= 15
    c.showPage()
    c.save()


def _chave(numero: int) -> str:
    return "276" + "666" + "45723174001192" + "55" + "001" + f"{numero:09d}" + "0" * 10


def _soma(valores) -> str:
    total = sum((Decimal(v.replace(".", "").replace(",", ".")) for v in valores), Decimal("0"))
    return f"{total:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _us(valor_br: str) -> str:
    # a Sefaz escreve em formato US (56.80; 128,450.60), o Dominio em BR
    return valor_br.replace(".", "X").replace(",", ".").replace("X", ",")


def _sefaz(caminho, notas, cnpj="12.345.678/0001-95"):
    linhas = [
        "SECRETARIA DE ESTADO DA FAZENDA DE ALAGOAS",
        f"CNPJ: {cnpj}",
        "Razão Social: FARMACIA MODELO LTDA",
        "Período: 01/09/2026 - 30/09/2026",
    ]
    for numero, valor in notas:
        linhas.append(
            f"05/09/2026 08:52:43 45.723.174/0011-92 FORNECEDOR TESTE LTDA 1 / {numero} "
            f"{_chave(numero)} Saída Ativa {_us(valor)}"
        )
    linhas += [f"Quantidade de Notas: {len(notas)}", f"Valor Total: {_us(_soma(v for _, v in notas))}"]
    _pdf(caminho, linhas)


def _dominio(caminho, notas, cnpj="12.345.678/0001-95", extra=()):
    linhas = [
        "ACOMPANHAMENTO DE ENTRADAS",
        "101 - FARMACIA MODELO LTDA",
        f"CNPJ: {cnpj}",
        "Período: 01/09/2026 até 30/09/2026",
    ]
    for numero, valor in notas:
        linhas.append(
            f"1 05/09/2026 05/09/2026 {numero} 1 36 6 FORNECEDOR TESTE LTDA "
            f"45.723.174/0011-92 240000015 1-102 101 AL {valor} ICMS 0,00"
        )
    _pdf(caminho, [*linhas, *extra, "Total Geral", f"{_soma(v for _, v in notas)}ICMS 0,00"])


@pytest.fixture
def entrada(tmp_path):
    pasta = tmp_path / "entrada"
    pasta.mkdir()
    return pasta


def _resumo(saida):
    aba = load_workbook(next(saida.rglob("*.xlsx")))["Resumo"]
    cab = [c.value for c in aba[1]]
    return [{cab[i]: c.value for i, c in enumerate(linha)} for linha in aba.iter_rows(min_row=2)]


def _log(saida):
    return next(saida.rglob("log_execucao.txt")).read_text(encoding="utf-8")


# --- testes com dados fabricados ---

def test_pasta_vazia_nao_quebra(tmp_path):
    entrada = tmp_path / "entrada"
    entrada.mkdir()
    saida = tmp_path / "saida"
    assert executa(entrada, saida) == 0


def test_pasta_de_entrada_inexistente_levanta(tmp_path):
    with pytest.raises(NotADirectoryError):
        executa(tmp_path / "nao_existe", tmp_path / "saida")


def test_empresa_divergente_gera_pdf_planilha_e_log(entrada, tmp_path):
    saida = tmp_path / "saida"
    _sefaz(entrada / "s.pdf", [(100, "50,00"), (101, "70,00")])
    _dominio(entrada / "d.pdf", [(100, "50,00")])

    assert executa(entrada, saida) == 1

    # subdiretório vem do período declarado, não da data de hoje
    pasta = saida / "2026-09"
    assert (pasta / "resumo_conferencia_2026-09.xlsx").is_file()
    assert [p.name for p in pasta.glob("101_*.pdf")] == ["101_FARMACIA_MODELO_LTDA.pdf"]
    [linha] = _resumo(saida)
    assert linha["Código"] == "101" and linha["Status"] == Status.DIVERGENTE
    assert linha[Categoria.FALTANDO] == 1
    assert linha["Período"] == "01/09/2026 a 30/09/2026"
    assert linha["Qtd Sefaz"] == 2 and linha["Qtd Domínio"] == 1
    assert Decimal(str(linha["Total Sefaz"])) == Decimal("120.00")
    log = _log(saida)
    assert "101" in log and "s.pdf" in log and "d.pdf" in log


def test_empresa_sem_ocorrencia_tambem_gera_pdf(entrada, tmp_path):
    saida = tmp_path / "saida"
    _sefaz(entrada / "s.pdf", [(100, "50,00")])
    _dominio(entrada / "d.pdf", [(100, "50,00")])
    assert executa(entrada, saida) == 0
    assert list((saida / "2026-09").glob("101_*.pdf"))
    assert _resumo(saida)[0]["Status"] == Status.OK


def test_agrega_linhas_de_cfop_da_mesma_nota_e_fecha_com_a_sefaz(entrada, tmp_path):
    # uma nota, duas linhas de CFOP no Domínio: 52,10 + 140,30 = 192,40 (valor da Sefaz)
    saida = tmp_path / "saida"
    _sefaz(entrada / "s.pdf", [(100, "192,40")])
    _dominio(entrada / "d.pdf", [(100, "52,10"), (100, "140,30")])

    assert executa(entrada, saida) == 0

    [linha] = _resumo(saida)
    assert linha["Status"] == Status.OK
    assert linha["Qtd Domínio"] == 1
    assert Decimal(str(linha["Total Domínio"])) == Decimal("192.40")
    assert Decimal(str(linha["Total Sefaz"])) == Decimal("192.40")
    assert linha[Categoria.VALOR_DIVERGENTE] == 0


def test_empresa_sem_notas_no_mes_com_relatorios_vazios_e_coerentes_sai_ok(entrada, tmp_path):
    saida = tmp_path / "saida"
    _sefaz(entrada / "s.pdf", [])
    _dominio(entrada / "d.pdf", [])
    assert executa(entrada, saida) == 0
    [linha] = _resumo(saida)
    assert linha["Status"] == Status.OK
    assert linha["Qtd Sefaz"] == 0 and linha["Qtd Domínio"] == 0


def test_sem_par_aparece_no_resumo_sem_pdf(entrada, tmp_path):
    saida = tmp_path / "saida"
    _sefaz(entrada / "s.pdf", [(100, "50,00")])
    assert executa(entrada, saida) == 0
    [linha] = _resumo(saida)
    assert linha["Status"] == Status.SEM_PAR and "dominio" in linha["Observação"].lower()
    assert not list(saida.rglob("*.pdf"))


def test_arquivo_ilegivel_nao_derruba_a_rodada(entrada, tmp_path):
    # spec §5.2: falha em uma empresa não interrompe as demais
    saida = tmp_path / "saida"
    _sefaz(entrada / "s.pdf", [(100, "50,00")])
    _dominio(entrada / "d.pdf", [(100, "50,00")])
    (entrada / "quebrado.pdf").write_bytes(b"isto nao e um PDF")

    assert executa(entrada, saida) == 0

    linhas = _resumo(saida)
    assert Status.OK in [l["Status"] for l in linhas]
    erro = next(l for l in linhas if l["Status"] == Status.ERRO_LEITURA)
    assert "quebrado.pdf" in erro["Observação"]
    assert "quebrado.pdf" in _log(saida)
    assert list(saida.rglob("101_*.pdf"))


def test_linha_fora_do_layout_derruba_so_a_empresa(entrada, tmp_path):
    # linha de dados sem CNPJ: pular em silêncio esconderia um FALTANDO
    saida = tmp_path / "saida"
    _sefaz(entrada / "s1.pdf", [(100, "50,00")])
    _dominio(entrada / "d1.pdf", [(100, "50,00")], extra=["1 05/09/2026 05/09/2026 101 1 36 lixo"])
    _sefaz(entrada / "s2.pdf", [(200, "10,00")], cnpj="11.222.333/0001-81")
    _dominio(entrada / "d2.pdf", [(200, "10,00")], cnpj="11.222.333/0001-81")

    assert executa(entrada, saida) == 0

    linhas = _resumo(saida)
    assert sorted(l["Status"] for l in linhas) == sorted([Status.OK, Status.ERRO_LEITURA])
    erro = next(l for l in linhas if l["Status"] == Status.ERRO_LEITURA)
    assert "d1.pdf" in erro["Observação"]
    assert "d1.pdf" in _log(saida)
    assert len(list(saida.rglob("*.pdf"))) == 1


def test_sefaz_com_nota_perdida_na_leitura_vira_erro_de_leitura_nao_ok(entrada, tmp_path):
    # chave partida em dois tokens: sem o rodape a empresa sairia OK com 0 notas lidas
    saida = tmp_path / "saida"
    _pdf(entrada / "s.pdf", [
        "SECRETARIA DE ESTADO DA FAZENDA DE ALAGOAS",
        "CNPJ: 12.345.678/0001-95",
        "Razão Social: FARMACIA MODELO LTDA",
        "Período: 01/09/2026 - 30/09/2026",
        f"05/09/2026 08:52:43 45.723.174/0011-92 FORNECEDOR TESTE LTDA 1 / 100 "
        f"{_chave(100)} Saída Ativa 50.00",
        f"05/09/2026 08:52:43 45.723.174/0011-92 FORNECEDOR TESTE LTDA 1 / 101 "
        f"{_chave(101)[:20]} {_chave(101)[20:]} Saída Ativa 70.00",
        "Quantidade de Notas: 2", "Valor Total: 120.00",
    ])
    _dominio(entrada / "d.pdf", [(100, "50,00"), (101, "70,00")])
    assert executa(entrada, saida) == 0
    [linha] = _resumo(saida)
    assert linha["Status"] == Status.ERRO_LEITURA
    assert "Quantidade de Notas" in linha["Observação"]


# --- teste de ouro: depende dos PDFs reais, locais e não versionados ---

def _pdfs_da_amostra() -> list[Path]:
    return sorted(AMOSTRA.glob("*.pdf")) if AMOSTRA.exists() else []


def _pares_da_amostra() -> list[Par]:
    """Pares Sefaz/Dominio da pasta, por CNPJ lido de dentro dos PDFs (nomes nao importam)."""
    if len(_pdfs_da_amostra()) < 2:
        return []
    return pareia([a for a in map(identifica, _pdfs_da_amostra()) if a])[0]


_PARES = _pares_da_amostra()


@pytest.fixture(scope="module")
def rodada_amostra(tmp_path_factory):
    pdfs = _pdfs_da_amostra()
    if len(pdfs) < 2:
        pytest.skip(
            f"a pasta {AMOSTRA} tem menos de dois PDFs: coloque nela os pares reais "
            "(relatorio da Sefaz e do Dominio de cada empresa)"
        )
    base = tmp_path_factory.mktemp("amostra")
    entrada, saida = base / "entrada", base / "saida"
    entrada.mkdir()
    for p in pdfs:
        shutil.copy(p, entrada)
    executa(entrada, saida)
    wb = load_workbook(next(saida.rglob("*.xlsx")))
    return wb, saida


@pytest.mark.parametrize("par", _PARES, ids=lambda p: f"empresa_{p.empresa.codigo}")
def test_teste_de_ouro_reconcilia_cada_empresa_com_os_proprios_rodapes(rodada_amostra, par):
    wb, _ = rodada_amostra
    # os leitores ja exigem que linhas lidas == rodape declarado (Sefaz) e Total Geral (Dominio)
    _, notas = le_sefaz(par.caminho_sefaz)
    _, lancamentos = le_dominio(par.caminho_dominio)
    qtd_declarada = len(notas)
    total_sefaz_declarado = sum((n.valor for n in notas), Decimal("0"))
    total_geral_declarado = sum((l.valor_contabil for l in lancamentos), Decimal("0"))

    aba = wb["Resumo"]
    cabecalho = [c.value for c in aba[1]]
    linhas = [dict(zip(cabecalho, r)) for r in aba.iter_rows(min_row=2, values_only=True)]
    linha = next(l for l in linhas if l["Código"] == par.empresa.codigo)

    assert linha["Status"] not in (Status.ERRO_LEITURA, Status.SEM_PAR), linha["Observação"]
    assert linha["Qtd Sefaz"] == qtd_declarada
    assert Decimal(str(linha["Total Sefaz"])) == total_sefaz_declarado

    # Total Dominio so soma a espécie NF-e; as demais viram ocorrencia NAO-NFe no Detalhe.
    det = wb["Detalhe"]
    cab_det = [c.value for c in det[1]]
    nao_nfe = sum(
        (Decimal(str(r[cab_det.index("Valor Domínio")])) for r in det.iter_rows(min_row=2, values_only=True)
         if r[cab_det.index("Código")] == par.empresa.codigo
         and r[cab_det.index("Categoria")] == Categoria.NAO_NFE),
        Decimal("0"),
    )
    assert Decimal(str(linha["Total Domínio"])) + nao_nfe == total_geral_declarado


@pytest.mark.parametrize("par", _PARES, ids=lambda p: f"empresa_{p.empresa.codigo}")
def test_gera_pdf_para_cada_empresa_conferida(rodada_amostra, par):
    _, saida = rodada_amostra
    assert [p.name for p in saida.rglob(f"{par.empresa.codigo}_*.pdf")]


def test_falha_ao_gravar_pdf_preserva_a_conferencia(entrada, tmp_path, monkeypatch):
    saida = tmp_path / "saida"
    _sefaz(entrada / "s.pdf", [(100, "50,00"), (101, "70,00")])
    _dominio(entrada / "d.pdf", [(100, "50,00")])

    def nega(*a, **k):
        raise PermissionError("pdf aberto")
    monkeypatch.setattr("main.escreve_relatorio", nega)

    assert executa(entrada, saida) == 1
    [linha] = _resumo(saida)
    assert linha["Status"] == Status.DIVERGENTE and linha[Categoria.FALTANDO] == 1
    assert "PDF" in linha["Observação"]
    assert "pdf aberto" in _log(saida)


def test_log_sobrevive_a_falha_da_planilha(entrada, tmp_path, monkeypatch):
    saida = tmp_path / "saida"
    _sefaz(entrada / "s.pdf", [(100, "50,00")])
    _dominio(entrada / "d.pdf", [(100, "50,00")])

    def nega(*a, **k):
        raise PermissionError("planilha aberta no Excel")
    monkeypatch.setattr("main.escreve_planilha", nega)

    with pytest.raises(PermissionError):
        executa(entrada, saida)
    assert "101" in _log(saida)


# --- aviso de mudanca de formato numerico chega ao log_execucao.txt e ao console ---

def test_sefaz_em_formato_brasileiro_avisa_no_log_e_no_console(entrada, tmp_path, capsys, monkeypatch):
    monkeypatch.setattr("tests.test_ponta_a_ponta._us", lambda v: v)  # Sefaz escrevendo em formato brasileiro
    _sefaz(entrada / "sefaz.pdf", [(1, "1.532,66"), (2, "10,00")])
    _dominio(entrada / "dominio.pdf", [(1, "1.532,66"), (2, "10,00")])
    saida = tmp_path / "saida"
    assert executa(entrada, saida) == 0
    log = _log(saida)
    assert "formato numérico do relatório da Sefaz mudou: esperado 1,234.56 e encontrado 1.234,56" in log
    assert "formato numérico do relatório da Sefaz mudou" in capsys.readouterr().out
    assert "relatório do Domínio mudou" not in log


def test_dominio_em_formato_americano_avisa_no_log(entrada, tmp_path):
    _sefaz(entrada / "sefaz.pdf", [(1, "1.532,66")])
    linhas = [
        "ACOMPANHAMENTO DE ENTRADAS", "101 - FARMACIA MODELO LTDA", "CNPJ: 12.345.678/0001-95",
        "Período: 01/09/2026 até 30/09/2026",
        "1 05/09/2026 05/09/2026 1 1 36 6 FORNECEDOR TESTE LTDA 45.723.174/0011-92 240000015 "
        "1-102 101 AL 1,532.66 ICMS 0.00",
        "Total Geral", "1,532.66ICMS 0.00",
    ]
    _pdf(entrada / "dominio.pdf", linhas)
    saida = tmp_path / "saida"
    assert executa(entrada, saida) == 0
    assert "formato numérico do relatório do Domínio mudou: esperado 1.234,56 e encontrado 1,234.56" in _log(saida)


def test_executa_nao_deixa_handler_no_logger_do_pacote(entrada, tmp_path):
    antes = list(logging.getLogger("conferencia").handlers)
    executa(entrada, tmp_path / "saida")
    assert logging.getLogger("conferencia").handlers == antes
    with pytest.raises(NotADirectoryError):
        executa(tmp_path / "nao_existe", tmp_path / "s2")
    assert logging.getLogger("conferencia").handlers == antes
