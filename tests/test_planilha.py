from decimal import Decimal

from openpyxl import load_workbook

from conferencia.modelo import Categoria, Status
from conferencia.planilha import escreve_planilha


def _linhas(aba):
    cab = [c.value for c in aba[1]]
    return [dict(zip(cab, [c.value for c in lin])) for lin in aba.iter_rows(min_row=2)]


def test_resumo_tem_uma_linha_por_empresa_inclusive_as_corretas(tmp_path, resultados_exemplo):
    destino = tmp_path / "r.xlsx"
    escreve_planilha(resultados_exemplo, destino)
    aba = load_workbook(destino)["Resumo"]
    assert aba.max_row == 1 + len(resultados_exemplo)


def test_detalhe_tem_uma_linha_por_ocorrencia(tmp_path, resultados_exemplo):
    destino = tmp_path / "r.xlsx"
    escreve_planilha(resultados_exemplo, destino)
    aba = load_workbook(destino)["Detalhe"]
    total = sum(len(r.ocorrencias) for r in resultados_exemplo)
    assert aba.max_row == 1 + total


def test_valores_sao_numericos_e_nao_texto(tmp_path, resultados_exemplo):
    destino = tmp_path / "r.xlsx"
    escreve_planilha(resultados_exemplo, destino)
    aba = load_workbook(destino)["Resumo"]
    cabecalho = [c.value for c in aba[1]]
    coluna = cabecalho.index("Total Sefaz") + 1
    celula = aba.cell(row=2, column=coluna)
    assert isinstance(celula.value, (int, float, Decimal))
    assert celula.number_format == "#,##0.00"


def test_empresa_sem_par_aparece_no_resumo_com_observacao(tmp_path, resultados_exemplo):
    destino = tmp_path / "r.xlsx"
    escreve_planilha(resultados_exemplo, destino)
    aba = load_workbook(destino)["Resumo"]
    status = [aba.cell(row=l, column=4).value for l in range(2, aba.max_row + 1)]
    assert Status.SEM_PAR in status
    sem_par = [l for l in _linhas(aba) if l["Status"] == Status.SEM_PAR][0]
    assert "99888777000155" in sem_par["Observação"]
    assert sem_par["Código"] is None


def test_planilha_vazia_nao_quebra(tmp_path):
    destino = tmp_path / "r.xlsx"
    escreve_planilha([], destino)
    assert load_workbook(destino)["Resumo"].max_row == 1


def test_resumo_separa_erros_de_alertas_e_conta_por_categoria(tmp_path, resultado_divergente):
    destino = tmp_path / "r.xlsx"
    escreve_planilha([resultado_divergente], destino)
    linha = _linhas(load_workbook(destino)["Resumo"])[0]
    assert linha["Erros"] == 1  # FALTANDO
    assert linha["Alertas"] == 1  # NÃO-NFe
    assert linha[Categoria.FALTANDO] == 1
    assert linha[Categoria.NAO_NFE] == 1
    assert linha[Categoria.SOBRANDO] == 0
    assert round(linha["Diferença"], 2) == 170.75


def test_detalhe_traz_valores_numericos_chave_e_lado_ausente_vazio(tmp_path, resultado_divergente):
    destino = tmp_path / "r.xlsx"
    escreve_planilha([resultado_divergente], destino)
    faltando, nao_nfe = _linhas(load_workbook(destino)["Detalhe"])
    assert faltando["Categoria"] == Categoria.FALTANDO
    assert faltando["Chave de Acesso"] == "2" * 44
    assert faltando["Valor Sefaz"] == 250.75
    assert faltando["Valor Domínio"] is None
    assert faltando["Diferença"] is None
    assert faltando["CNPJ Empresa"] == "44555666000172"
    assert nao_nfe["Chave de Acesso"] is None
    assert nao_nfe["Valor Domínio"] == 80.0


def test_detalhe_calcula_diferenca_quando_ha_os_dois_valores(tmp_path, resultado_divergente):
    from dataclasses import replace

    oc = replace(
        resultado_divergente.ocorrencias[0],
        categoria=Categoria.VALOR_DIVERGENTE,
        valor_dominio=Decimal("200.00"),
    )
    resultado_divergente.ocorrencias = [oc]
    destino = tmp_path / "r.xlsx"
    escreve_planilha([resultado_divergente], destino)
    assert _linhas(load_workbook(destino)["Detalhe"])[0]["Diferença"] == 50.75


def test_primeira_linha_congelada_e_autofiltro_nas_duas_abas(tmp_path, resultados_exemplo):
    destino = tmp_path / "r.xlsx"
    escreve_planilha(resultados_exemplo, destino)
    wb = load_workbook(destino)
    for nome in ("Resumo", "Detalhe"):
        assert wb[nome].freeze_panes == "A2"
        assert wb[nome].auto_filter.ref is not None
