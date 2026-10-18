from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from conferencia.modelo import CATEGORIAS_ERRO, Categoria, ResultadoEmpresa

FORMATO_MOEDA = "#,##0.00"
FORMATO_DATA = "dd/mm/yyyy"
CENTAVO = Decimal("0.01")

CATEGORIAS = [
    valor for nome, valor in vars(Categoria).items() if not nome.startswith("_")
]

CABECALHO_RESUMO = [
    "Código", "CNPJ", "Razão Social", "Status", "Erros", "Alertas", *CATEGORIAS,
    "Qtd Sefaz", "Total Sefaz", "Qtd Domínio", "Total Domínio", "Diferença", "Período",
    "Observação",
]
CABECALHO_DETALHE = [
    "Código", "CNPJ Empresa", "Categoria", "CNPJ Emitente", "Fornecedor", "Série",
    "Número", "Data Emissão", "Valor Sefaz", "Valor Domínio", "Diferença",
    "Chave de Acesso",
]


def _dinheiro(valor: Decimal | None) -> float | None:
    """Arredonda no Decimal e só então converte; None vira célula vazia."""
    if valor is None:
        return None
    return float(valor.quantize(CENTAVO, rounding=ROUND_HALF_UP))


def _diferenca(sefaz: Decimal | None, dominio: Decimal | None) -> float | None:
    if sefaz is None or dominio is None:
        return None
    return _dinheiro(sefaz - dominio)


def _prepara_aba(aba: Worksheet, cabecalho: list[str]) -> None:
    aba.append(cabecalho)
    aba.freeze_panes = "A2"


def _fecha_aba(aba: Worksheet, formatos: dict[str, str]) -> None:
    """Aplica o formato por nome de coluna e liga o autofiltro."""
    cabecalho = [c.value for c in aba[1]]
    for nome, formato in formatos.items():
        coluna = cabecalho.index(nome) + 1
        for linha in aba.iter_rows(min_row=2, min_col=coluna, max_col=coluna):
            linha[0].number_format = formato
    aba.auto_filter.ref = aba.dimensions


def escreve_planilha(resultados: list[ResultadoEmpresa], destino: Path) -> None:
    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    detalhe = wb.create_sheet("Detalhe")
    _prepara_aba(resumo, CABECALHO_RESUMO)
    _prepara_aba(detalhe, CABECALHO_DETALHE)

    for r in resultados:
        codigo = r.empresa.codigo if r.empresa else None
        cnpj = r.empresa.cnpj if r.empresa else None
        razao = r.empresa.razao_social if r.empresa else None
        periodo = (
            f"{r.empresa.periodo_inicio:%d/%m/%Y} a {r.empresa.periodo_fim:%d/%m/%Y}"
            if r.empresa else None
        )
        erros = sum(1 for o in r.ocorrencias if o.categoria in CATEGORIAS_ERRO)
        por_categoria = [
            sum(1 for o in r.ocorrencias if o.categoria == c) for c in CATEGORIAS
        ]
        resumo.append([
            codigo, cnpj, razao, r.status, erros, len(r.ocorrencias) - erros,
            *por_categoria,
            r.qtd_sefaz, _dinheiro(r.total_sefaz), r.qtd_dominio,
            _dinheiro(r.total_dominio), _diferenca(r.total_sefaz, r.total_dominio),
            periodo, r.erro,
        ])
        for o in r.ocorrencias:
            detalhe.append([
                codigo, cnpj, o.categoria, o.cnpj_emitente, o.nome_emitente, o.serie,
                o.numero, o.data_emissao, _dinheiro(o.valor_sefaz),
                _dinheiro(o.valor_dominio), _diferenca(o.valor_sefaz, o.valor_dominio),
                o.chave,
            ])

    _fecha_aba(resumo, dict.fromkeys(["Total Sefaz", "Total Domínio", "Diferença"], FORMATO_MOEDA))
    _fecha_aba(
        detalhe,
        {
            **dict.fromkeys(["Valor Sefaz", "Valor Domínio", "Diferença"], FORMATO_MOEDA),
            "Data Emissão": FORMATO_DATA,
        },
    )
    wb.save(destino)
